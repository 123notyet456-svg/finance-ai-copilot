from io import BytesIO

from docx import Document
import pytest

from finance_copilot.core.document_processing import parse_uploaded_document
from finance_copilot.core.exceptions import FileValidationError


def test_parse_docx_includes_paragraphs_and_tables() -> None:
    document = Document()
    document.add_heading("差旅制度", level=1)
    document.add_paragraph("上海住宿标准为 600 元。")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "城市"
    table.rows[0].cells[1].text = "上海"
    buffer = BytesIO()
    document.save(buffer)

    parsed = parse_uploaded_document("travel.docx", buffer.getvalue())
    assert "上海住宿标准" in parsed.content
    assert "城市 | 上海" in parsed.content
    assert parsed.file_type == "docx"


def test_reject_unsupported_upload() -> None:
    with pytest.raises(FileValidationError):
        parse_uploaded_document("policy.exe", b"unsafe")
