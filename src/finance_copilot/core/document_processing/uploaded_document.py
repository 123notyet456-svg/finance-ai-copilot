"""Safe text extraction for uploaded policy documents."""

from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
from pathlib import Path

from finance_copilot.core.exceptions import DocumentProcessingError, FileValidationError


ALLOWED_UPLOAD_SUFFIXES = frozenset({".txt", ".md", ".markdown", ".pdf", ".docx"})
MAX_UPLOAD_BYTES = 10 * 1024 * 1024


@dataclass(frozen=True, slots=True)
class ParsedUpload:
    file_name: str
    file_type: str
    content: str
    content_hash: str


def parse_uploaded_document(file_name: str, data: bytes) -> ParsedUpload:
    safe_name = Path(file_name).name
    suffix = Path(safe_name).suffix.lower()
    if suffix not in ALLOWED_UPLOAD_SUFFIXES:
        raise FileValidationError(f"不支持的文件类型：{suffix or '无扩展名'}")
    if not data:
        raise FileValidationError("文件不能为空")
    if len(data) > MAX_UPLOAD_BYTES:
        raise FileValidationError(f"文件超过 10 MB 限制：{len(data)} bytes")

    try:
        if suffix in {".txt", ".md", ".markdown"}:
            content = data.decode("utf-8-sig").strip()
        elif suffix == ".pdf":
            content = _parse_pdf(data)
        else:
            content = _parse_docx(data)
    except UnicodeDecodeError as exc:
        raise DocumentProcessingError("文本文件必须使用 UTF-8 编码") from exc

    if not content:
        raise DocumentProcessingError("文档未提取到可索引文本")
    return ParsedUpload(
        file_name=safe_name,
        file_type=suffix.removeprefix("."),
        content=content,
        content_hash=sha256(data).hexdigest(),
    )


def _parse_pdf(data: bytes) -> str:
    try:
        import fitz  # type: ignore[import-not-found]

        with fitz.open(stream=data, filetype="pdf") as document:
            text = "\n".join(page.get_text() for page in document).strip()
    except Exception as exc:
        raise DocumentProcessingError("PDF 解析失败") from exc
    if not text:
        raise DocumentProcessingError("PDF 未提取到文本；扫描件 OCR 尚未开放")
    return text


def _parse_docx(data: bytes) -> str:
    try:
        from docx import Document  # type: ignore[import-not-found]

        document = Document(BytesIO(data))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
        table_rows = [
            " | ".join(cell.text.strip() for cell in row.cells)
            for table in document.tables
            for row in table.rows
        ]
        return "\n".join(item for item in [*paragraphs, *table_rows] if item).strip()
    except Exception as exc:
        raise DocumentProcessingError("DOCX 解析失败") from exc
